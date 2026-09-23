"""Fill the official SIH 2026 idea template with the ThermalSentinel content.

Works on the real template file, so the theme, graphics and official pointer wording are kept.
Removes the "@SIH Idea submission- Template" footer on every slide and the instructions slide.
"""
import copy
import json
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(HERE)
M = json.load(open(os.path.join(DOCS, 'metrics.json'), encoding='utf-8'))

TEAM = 'Fire Orbit'
TEAM_ID = ''
DEMO = 'https://rahulskandagal.github.io/thermal-sentinel/'
REPO = 'https://github.com/rahulskandagal/thermal-sentinel'

NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x1F, 0x29, 0x37)
MUT = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC8, 0xD0, 0xE0)
SOFT = RGBColor(0xF4, 0xF6, 0xFB)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
ORANGE = RGBColor(0xE8, 0x59, 0x0C)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xB9, 0x1C, 0x1C)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
LINKC = RGBColor(0x1A, 0x0D, 0xAB)
SANS = 'Arial'

prs = Presentation(os.path.join(HERE, 'template.pptx'))
S = prs.slides


# ---------------------------------------------------------------- helpers
def drop(shape):
    shape._element.getparent().remove(shape._element)


def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    return None


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    return box, tf


def para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def run(p, text, size=10, bold=False, color=INK, italic=False, font=SANS, underline=False, link=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    f.underline = underline
    if link:
        r.hyperlink.address = link
        r.font.color.rgb = LINKC
        r.font.underline = True
    return r


def text(slide, x, y, w, h, s, size=10, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False,
         anchor=MSO_ANCHOR.TOP, font=SANS, space_after=0):
    box, tf = tb(slide, x, y, w, h, anchor)
    for i, line in enumerate(str(s).split('\n')):
        p = para(tf, i == 0)
        p.alignment = align
        p.space_after = Pt(space_after)
        run(p, line, size, bold, color, italic, font)
    return box


def bullets(slide, x, y, w, h, items, size=10, color=INK, space=4, bullet_char='•  '):
    box, tf = tb(slide, x, y, w, h)
    for i, item in enumerate(items):
        p = para(tf, i == 0)
        p.space_after = Pt(space)
        if isinstance(item, tuple):
            run(p, bullet_char + item[0] + ' ', size, True, color)
            run(p, item[1], size, False, INK)
        else:
            run(p, bullet_char + item, size, False, color)
    return box


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    if adj is not None and sh.adjustments:
        sh.adjustments[0] = adj
    return sh


def card(slide, x, y, w, h, title, items, accent, size=9.5, tsize=11):
    rect(slide, x, y, w, h, SOFT, LINE, 1.0, adj=0.04)
    rect(slide, x, y, 0.07, h, accent, None, adj=0.4)
    text(slide, x + 0.2, y + 0.11, w - 0.35, 0.3, title, tsize, True, accent)
    bullets(slide, x + 0.2, y + 0.47, w - 0.38, h - 0.6, items, size)


def chip(slide, x, y, w, h, label, value, accent):
    rect(slide, x, y, w, h, accent, None, adj=0.12)
    text(slide, x + 0.08, y + 0.06, w - 0.16, 0.34, value, 16, True, WHITE, PP_ALIGN.CENTER)
    text(slide, x + 0.08, y + 0.42, w - 0.16, 0.34, label, 8, False, WHITE, PP_ALIGN.CENTER)


def numbered(slide, x, y, d, n, fill, color=WHITE):
    c = rect(slide, x, y, d, d, fill, None, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, str(n), 11, True, color)
    return c


def arrow(slide, x, y, w, h=0.0):
    ln = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.14))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln


def pic(slide, name, x, y, w, h=None):
    return slide.shapes.add_picture(os.path.join(DOCS, name), Inches(x), Inches(y), Inches(w), Inches(h) if h else None)


# ---------------------------------------------------------------- global cleanup
for s in S:
    for sh in list(s.shapes):
        # remove the "@SIH Idea submission- Template" footer text
        if sh.has_text_frame and '@SIH Idea submission' in sh.text_frame.text:
            drop(sh)
        # team name oval
        elif sh.has_text_frame and sh.text_frame.text.strip() == 'Your Team Name':
            tf = sh.text_frame
            tf.paragraphs[0].runs[0].text = TEAM

# ---------------------------------------------------------------- slide 1: title page
s1 = S[0]
box = by_id(s1, 10)
tf = box.text_frame
rows = [('Problem Statement ID - ', 'SIH26162'),
        ('Problem Statement Title- ', 'AI-Based Detection and Classification of Industrial Fires and Persistent Thermal Sources Using NASA FIRMS, OSM & Satellite Data'),
        ('Theme- ', 'Disaster Management'),
        ('PS Category- ', 'Software'),
        ('Team ID- ', TEAM_ID),
        ('Team Name (Registered on portal) - ', TEAM)]
# keep the first (empty) paragraph, rewrite the rest
for p in list(tf.paragraphs)[1:]:
    p._element.getparent().remove(p._element)
for i, (k, v) in enumerate(rows):
    p = tf.add_paragraph()
    p.space_after = Pt(10)
    run(p, k, 17, True, RGBColor(0, 0, 0))
    run(p, v, 17, False, RGBColor(0, 0, 0))
text(s1, 0.36, 6.35, 6.6, 0.4, 'Organisation: National Technical Research Organisation (NTRO)', 13, False, MUT, italic=True)

# ---------------------------------------------------------------- slide 2: proposed solution
s2 = S[1]
title2 = by_id(s2, 15361)
title2.top, title2.height = Inches(0.12), Inches(0.95)
t2 = title2.text_frame
t2.word_wrap = True
for _p in list(t2.paragraphs)[1:]:
    _p._element.getparent().remove(_p._element)
_p0 = t2.paragraphs[0]
for _r in list(_p0.runs):
    _r._r.getparent().remove(_r._r)
run(_p0, 'IDEA TITLE: ThermalSentinel', 30, True, RGBColor(0, 0, 0))
_p0.alignment = PP_ALIGN.CENTER

ptr = by_id(s2, 15362)
ptr.left, ptr.top, ptr.width, ptr.height = Inches(0.36), Inches(1.2), Inches(12.6), Inches(0.42)
ptf = ptr.text_frame
for p in list(ptf.paragraphs)[1:]:
    p._element.getparent().remove(p._element)
p0 = ptf.paragraphs[0]
for r in list(p0.runs)[1:]:
    r._r.getparent().remove(r._r)
p0.runs[0].text = 'Proposed Solution (Describe your Idea/Solution/Prototype)'
p0.runs[0].font.size = Pt(20)

card(s2, 0.36, 1.75, 4.0, 2.55, 'Detailed explanation of the proposed solution', [
    'NASA FIRMS publishes every VIIRS and MODIS thermal anomaly within about three hours, but only as a hot pixel with no meaning.',
    'ThermalSentinel groups detections into thermal sources, tracks each source for 90 days, and adds context from OpenStreetMap.',
    'A hybrid classifier labels every detection as industrial fire, gas flare, mining or coal-seam fire, agricultural burn, wildfire or other.',
    'An alert is raised only when a known source burns far above its own baseline.'], BLUE, 10)

card(s2, 4.56, 1.75, 4.0, 2.55, 'How it addresses the problem', [
    'Industrial fires are segregated from forest fires and other natural fires, which is the core ask of the problem statement.',
    'Results are stored in a spatial database and shown as a map overlay, with GeoJSON export for QGIS and ArcGIS.',
    'Only 138 of 22,179 detections reach an operator, so a control room reviews the events that matter.',
    'Every label carries plain reasons, so an analyst can confirm or override it.'], GREEN, 10)

card(s2, 8.76, 1.75, 4.2, 2.55, 'Innovation and uniqueness of the solution', [
    'Classification of the source over 90 days instead of a single pixel.',
    'Alerts measured against each site\'s own normal level, not a fixed threshold.',
    'Fusion of OpenStreetMap infrastructure and land cover with the thermal signal.',
    'Persistent heat with no mapped facility is itself reported as a finding.',
    'Built only on free and open data, and already running as a public prototype.'], PURPLE, 10)

for i, (v, l, c) in enumerate([('22,179', 'detections classified', BLUE), ('114', 'persistent sources', TEAL),
                               ('138', 'incident alerts', RED), ('6', 'thermal classes', ORANGE),
                               ('99.9 %', 'accuracy on the labelled archive', GREEN)]):
    chip(s2, 0.36 + i * 1.72, 4.5, 1.55, 0.82, l, v, c)

pic(s2, 'dashboard.png', 9.2, 4.45, 3.76)
text(s2, 9.2, 6.55, 3.76, 0.3, 'Prototype GIS dashboard, live at ' + DEMO.replace('https://', ''), 8, False, MUT, PP_ALIGN.CENTER, italic=True)
text(s2, 0.36, 5.5, 8.6, 0.9,
     'Working prototype: the whole 90-day archive is classified in about seven seconds on a laptop, and the dashboard, '
     'the GeoJSON API and the QGIS export are already public. One environment variable switches the same pipeline to live NASA FIRMS data.',
     11, False, INK)

# ---------------------------------------------------------------- slide 3: technical approach
s3 = S[2]
ptr = by_id(s3, 17410)
ptr.left, ptr.top, ptr.width, ptr.height = Inches(0.36), Inches(1.18), Inches(6.5), Inches(0.5)
ptf = ptr.text_frame
ptf.word_wrap = True
for p in list(ptf.paragraphs)[1:]:
    p._element.getparent().remove(p._element)
p0 = ptf.paragraphs[0]
for r in list(p0.runs)[1:]:
    r._r.getparent().remove(r._r)
p0.runs[0].text = 'Technologies to be used (e.g. programming languages, frameworks, hardware)'
p0.runs[0].font.size = Pt(13)

tech = [('Data', 'NASA FIRMS Area API (VIIRS 375 m, MODIS 1 km), OpenStreetMap Overpass, Copernicus Sentinel-2, NASA Worldview', BLUE),
        ('AI and ML', 'Python, scikit-learn (haversine DBSCAN, HistGradientBoosting), rule engine, shapely STRtree', PURPLE),
        ('Backend', 'FastAPI REST serving GeoJSON, SQLite with a PostGIS-ready schema, OGC-standard exports', TEAL),
        ('Frontend', 'React 18 with Vite, Leaflet canvas renderer for 40,000 points, QGIS and ArcGIS export', ORANGE),
        ('Hardware', 'None. Software only: a laptop for the demo, a single Linux VM, NIC cloud or an air-gapped server', GREEN)]
for i, (k, v, c) in enumerate(tech):
    y = 1.78 + i * 0.42
    rect(s3, 0.36, y, 1.3, 0.32, c, None, adj=0.18)
    text(s3, 0.36, y + 0.03, 1.3, 0.26, k, 9.5, True, WHITE, PP_ALIGN.CENTER)
    text(s3, 1.76, y + 0.03, 5.1, 0.36, v, 9, False, INK)

text(s3, 7.1, 1.18, 5.85, 0.5, 'Methodology and process for implementation (Flow Charts/Images/ working prototype)', 13, True, NAVY)
steps = [('Ingest', 'FIRMS VIIRS and MODIS pulled in 10-day chunks, both sensors normalised, overlapping passes removed', BLUE),
         ('Cluster', 'Haversine DBSCAN at 750 m turns pixel-jittered detections into thermal sources (22,179 into 4,419)', PURPLE),
         ('Track', '90-day history per source: active days, span, night fraction, FRP mean, variation and z-score', TEAL),
         ('Contextualise', 'OpenStreetMap facility type and distance, plus land cover by point-in-polygon lookup', ORANGE),
         ('Classify', 'Explainable rules and gradient-boosted trees vote over 34 features and return reasons', RED),
         ('Alert and serve', 'FRP at 2.5 times the source mean raises an incident; results served as GeoJSON to the GIS dashboard', GREEN)]
for i, (k, v, c) in enumerate(steps):
    y = 1.78 + i * 0.58
    numbered(s3, 7.1, y, 0.32, i + 1, c)
    text(s3, 7.54, y, 5.4, 0.22, k, 10, True, c)
    text(s3, 7.54, y + 0.21, 5.4, 0.34, v, 8.5, False, INK)

rect(s3, 0.36, 3.95, 6.5, 0.86, SOFT, LINE, 1.0, adj=0.05)
box, tf = tb(s3, 0.5, 4.03, 6.24, 0.72)
p = tf.paragraphs[0]
run(p, 'Core of the model: ', 9.5, True, NAVY)
run(p, '34 features covering radiometry, persistence, OpenStreetMap context, land cover and Indian burn-season priors. '
       'Rules give auditable priors, the gradient-boosted model resolves the ambiguous cases, and agreement between them raises confidence. '
       'Training and testing use a source-aware split so no thermal source appears on both sides.', 9.5, False, INK)

pic(s3, 'dashboard.png', 0.36, 4.95, 3.18, 1.79)
pic(s3, 'ui_jharia.png', 3.68, 4.95, 3.18, 1.79)
text(s3, 0.36, 6.76, 6.5, 0.24,
     'Prototype: all-India classification, and a persistent coal-seam source at Jharia with its 90-day history.',
     8, False, MUT, PP_ALIGN.CENTER, italic=True)
rect(s3, 7.1, 5.3, 5.85, 1.35, SOFT, LINE, 1.0, adj=0.05)
text(s3, 7.26, 5.4, 5.5, 0.28, 'Prototype and source code', 10.5, True, NAVY)
box, tf = tb(s3, 7.26, 5.7, 5.5, 0.9)
p = tf.paragraphs[0]
run(p, 'Live demo: ', 9, True, INK)
run(p, DEMO, 8.5, link=DEMO)
p = tf.add_paragraph()
run(p, 'Code: ', 9, True, INK)
run(p, REPO, 8.5, link=REPO)

# ---------------------------------------------------------------- slide 4: feasibility and viability
s4 = S[3]
ptr = by_id(s4, 17410)
drop(ptr)
cards = [('Analysis of the feasibility of the idea', [
    ('Data is available: ', 'the FIRMS key is issued instantly and OpenStreetMap and Copernicus are open, so there is no procurement.'),
    ('It already runs: ', 'the full 90-day archive of 22,179 detections is classified in about seven seconds on a laptop.'),
    ('Light to deploy: ', 'FastAPI, SQLite and React run on a single VM, NIC cloud or an air-gapped server.'),
    ('Live in one step: ', 'a single environment variable moves the same pipeline to near-real-time FIRMS data.')], GREEN),
    ('Potential challenges and risks', [
        ('Thermal is not fire: ', 'steady process heat and flares look like fires in the raw feed.'),
        ('Incomplete mapping: ', 'OpenStreetMap misses many small Indian industrial units.'),
        ('Data gaps: ', 'cloud cover and pixel jitter break single-pass observations.'),
        ('Few labelled events: ', 'confirmed incident records for supervised learning are scarce.'),
        ('Trust: ', 'agencies will not act on a label they cannot question.')], RED),
    ('Strategies for overcoming these challenges', [
        ('Persistence and context: ', '90-day behaviour plus facility and land-cover context separate routine heat from fire.'),
        ('Baseline alerts: ', 'each site is compared with its own normal, so flares stay silent until they spike.'),
        ('Multi-satellite window: ', 'three VIIRS satellites and a 90-day window absorb cloud gaps and jitter.'),
        ('Feedback loop: ', 'analyst confirmations become labels and retrain the model.'),
        ('Explainability: ', 'every decision lists its reasons and can be overridden.')], BLUE)]
for i, (title, items, accent) in enumerate(cards):
    x = 0.36 + i * 4.24
    rect(s4, x, 1.2, 4.08, 3.6, SOFT, LINE, 1.0, adj=0.04)
    rect(s4, x, 1.2, 0.07, 3.6, accent, None, adj=0.4)
    text(s4, x + 0.2, 1.3, 3.75, 0.5, title, 13, True, accent)
    box, tf = tb(s4, x + 0.2, 1.88, 3.7, 2.8)
    for j, (k, v) in enumerate(items):
        p = para(tf, j == 0)
        p.space_after = Pt(6)
        run(p, '•  ' + k, 10, True, INK)
        run(p, v, 10, False, INK)

text(s4, 0.36, 4.95, 12.6, 0.3, 'Viability: what it costs and how it scales', 13, True, NAVY)
facts = [('Rs 0', 'data licensing, all sources open'), ('Rs 3,000 - 5,000', 'per month for one cloud VM'),
         ('7 seconds', 'to classify a 90-day national archive'), ('India-wide', 'OSM queried only around hotspot cells'),
         ('GeoJSON', 'plugs into Bhuvan, NDEM, QGIS')]
for i, (v, l) in enumerate(facts):
    x = 0.36 + i * 2.55
    rect(s4, x, 5.3, 2.4, 0.82, SOFT, LINE, 1.0, adj=0.06)
    text(s4, x + 0.1, 5.38, 2.2, 0.3, v, 13, True, NAVY, PP_ALIGN.CENTER)
    text(s4, x + 0.1, 5.7, 2.2, 0.36, l, 8.5, False, MUT, PP_ALIGN.CENTER)
text(s4, 0.36, 6.25, 12.6, 0.5,
     'Roadmap: the idea and the prototype are done and public. The next four weeks add live FIRMS ingestion for all India and validation '
     'against documented incidents, followed by alerting and a Sentinel-2 verification model, and then integration with agency GIS systems.',
     10.5, False, INK)

# ---------------------------------------------------------------- slide 5: impact and benefits
s5 = S[4]
ptr = by_id(s5, 17410)
drop(ptr)
text(s5, 0.36, 1.2, 6.2, 0.36, 'Potential impact on the target audience', 15, True, NAVY)
rows = [('Disaster management (NDMA, SDMAs, fire services)', 'Alerts only on abnormal industrial heat, with the facility name and the evidence, so response reaches the right plant sooner.'),
        ('NTRO and geospatial analysts', 'A national registry of persistent thermal sources with 90 days of activity history, including sites that are not on any map.'),
        ('CPCB, MoEFCC and state boards', 'Fires attributed to stubble burning, flaring or industry, district by district, for air-quality action.'),
        ('Industry and insurers', 'A verifiable incident timeline for any site, with satellite imagery links for claims and compliance.')]
for i, (who, what) in enumerate(rows):
    y = 1.62 + i * 0.86
    rect(s5, 0.36, y, 6.2, 0.78, SOFT, LINE, 1.0, adj=0.06)
    numbered(s5, 0.5, y + 0.2, 0.34, i + 1, NAVY)
    text(s5, 0.96, y + 0.07, 5.45, 0.26, who, 10.5, True, NAVY)
    text(s5, 0.96, y + 0.32, 5.45, 0.42, what, 9.5, False, INK)

text(s5, 6.95, 1.2, 6.0, 0.36, 'Benefits of the solution (social, economic, environmental, etc.)', 15, True, NAVY)
bens = [('Social', 'Communities living beside refineries, steel plants and chemical complexes are warned earlier, and fire services are used where they are needed.', GREEN),
        ('Economic', 'No data licensing cost and a single VM to run. Early detection reduces losses, and the open stack avoids paid geospatial platforms.', ORANGE),
        ('Environmental', 'Agricultural burning is separated from industrial emissions, and flaring and coal-seam fires are inventoried for climate reporting.', TEAL),
        ('Operational', 'A control room reviews 138 events instead of 22,179, and every decision is explainable and auditable.', PURPLE)]
for i, (k, v, c) in enumerate(bens):
    y = 1.62 + i * 0.86
    rect(s5, 6.95, y, 6.0, 0.78, SOFT, LINE, 1.0, adj=0.06)
    rect(s5, 6.95, y, 0.07, 0.78, c, None, adj=0.4)
    text(s5, 7.15, y + 0.07, 5.6, 0.26, k, 10.5, True, c)
    text(s5, 7.15, y + 0.32, 5.6, 0.42, v, 9.5, False, INK)

text(s5, 0.36, 5.2, 12.6, 0.3, 'Measured on the labelled 90-day archive', 13, True, NAVY)
kpi = [('99 %', 'fewer hotspots to triage'), ('22,179', 'detections classified'), ('114', 'persistent sources registered'),
       ('138', 'incident candidates'), ('6', 'classes separated')]
for i, (v, l) in enumerate(kpi):
    x = 0.36 + i * 2.55
    rect(s5, x, 5.55, 2.4, 0.82, NAVY, None, adj=0.06)
    text(s5, x + 0.1, 5.63, 2.2, 0.3, v, 15, True, WHITE, PP_ALIGN.CENTER)
    text(s5, x + 0.1, 5.96, 2.2, 0.34, l, 8.5, False, RGBColor(0xD6, 0xDE, 0xF0), PP_ALIGN.CENTER)
text(s5, 0.36, 6.5, 12.6, 0.3,
     'Example alerts found: Bhilai Steel at 414 MW against a 32 MW baseline, Jamnagar Refinery at 207 MW against 19 MW, Tata Steel at 101 MW against 26 MW.',
     10, False, INK, PP_ALIGN.CENTER)

# ---------------------------------------------------------------- slide 6: research and references
s6 = S[5]
ptr = by_id(s6, 17410)
ptr.left, ptr.top, ptr.width, ptr.height = Inches(0.36), Inches(1.2), Inches(12.6), Inches(0.36)
ptf = ptr.text_frame
for p in list(ptf.paragraphs)[1:]:
    p._element.getparent().remove(p._element)
p0 = ptf.paragraphs[0]
for r in list(p0.runs)[1:]:
    r._r.getparent().remove(r._r)
p0.runs[0].text = 'Details / Links of the reference and research work'
p0.runs[0].font.size = Pt(15)

refs = [('Problem statement', 'NTRO, PS SIH26162, Smart India Hackathon 2026: segregate industrial fires from forest and other natural fires, with GIS storage and map-overlay visualisation.', 'https://sih.gov.in'),
        ('Hotspot data', 'NASA FIRMS Area API: VIIRS 375 m (Suomi-NPP, NOAA-20, NOAA-21) and MODIS 1 km active-fire products, near real time and archive.', 'https://firms.modaps.eosdis.nasa.gov/api/area/'),
        ('Sensor science', 'Schroeder, W. et al. (2014). The new VIIRS 375 m active fire detection data product. Remote Sensing of Environment, 143, 85-96.', 'https://doi.org/10.1016/j.rse.2013.12.008'),
        ('Flare signatures', 'Elvidge, C. D. et al. (2013). VIIRS Nightfire: satellite pyrometry at night. Remote Sensing, 5(9), 4423-4449.', 'https://doi.org/10.3390/rs5094423'),
        ('Industrial heat', 'Liu, Y. et al. (2018). Identifying industrial heat sources using time series of the VIIRS Nightfire product. Remote Sensing of Environment, 204.', 'https://www.sciencedirect.com/journal/remote-sensing-of-environment'),
        ('Context data', 'OpenStreetMap Overpass API and tagging: industrial, man_made=flare or works, power=plant, landuse=quarry, farmland or forest.', 'https://wiki.openstreetmap.org/wiki/Overpass_API'),
        ('Verification', 'Copernicus Data Space, Sentinel-2 SWIR bands B11 and B12 for visual confirmation, and ESA WorldCover 10 m land cover.', 'https://dataspace.copernicus.eu'),
        ('Seasonal priors', 'ICAR-IARI CREAMS and CPCB crop-residue burning bulletins, used as the October-November and April-May seasonality prior.', 'https://creams.iari.res.in'),
        ('Methods', 'Ester, M. et al. (1996). DBSCAN, KDD-96, used for source grouping; scikit-learn HistGradientBoostingClassifier documentation.', 'https://scikit-learn.org')]
for i, (tag, body, url) in enumerate(refs):
    col, row = i // 5, i % 5
    x = 0.36 + col * 6.4
    y = 1.68 + row * 0.72
    numbered(s6, x, y + 0.06, 0.3, i + 1, NAVY)
    text(s6, x + 0.42, y, 5.7, 0.22, tag, 10, True, NAVY)
    text(s6, x + 0.42, y + 0.21, 5.7, 0.34, body, 8.5, False, INK)
    box, tf = tb(s6, x + 0.42, y + 0.52, 5.7, 0.18)
    run(tf.paragraphs[0], url, 8, link=url)

rect(s6, 6.76, 5.3, 6.2, 1.5, SOFT, LINE, 1.0, adj=0.05)
text(s6, 6.92, 5.4, 5.9, 0.28, 'Our own work', 10.5, True, NAVY)
box, tf = tb(s6, 6.92, 5.7, 5.9, 1.0)
p = tf.paragraphs[0]
run(p, 'Working prototype: ', 9.5, True, INK)
run(p, DEMO, 9, link=DEMO)
p = tf.add_paragraph()
p.space_before = Pt(4)
run(p, 'Source code and documentation: ', 9.5, True, INK)
run(p, REPO, 9, link=REPO)
pic(s6, 'dashboard.png', 0.36, 5.3, 2.66, 1.5)
pic(s6, 'ui_punjab.png', 3.2, 5.3, 2.66, 1.5)
text(s6, 0.36, 6.76, 5.5, 0.2, 'Prototype: all-India view and the Punjab crop-burn season', 8, False, MUT, PP_ALIGN.CENTER, italic=True)

# ---------------------------------------------------------------- remove the instructions slide
xml_slides = prs.slides._sldIdLst
slide_ids = list(xml_slides)
last = slide_ids[-1]
prs.part.drop_rel(last.rId)
xml_slides.remove(last)

# ---------------------------------------------------------------- properties
cp = prs.core_properties
cp.title = 'SIH26162 ThermalSentinel - Team Fire Orbit'
cp.subject = 'Smart India Hackathon 2026 idea submission (PS SIH26162, NTRO)'
cp.author = 'Team Fire Orbit'
cp.last_modified_by = 'Team Fire Orbit'
cp.keywords = 'SIH 2026; SIH26162; NTRO; thermal anomalies; NASA FIRMS'
cp.comments = ''
cp.category = ''

out = os.path.join(HERE, 'FireOrbit_SIH26162_ThermalSentinel_OfficialTemplate.pptx')
prs.save(out)
print('wrote', out, '| slides:', len(prs.slides.__iter__.__self__._sldIdLst))
